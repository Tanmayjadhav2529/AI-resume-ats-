import importlib
import io
import unittest
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from docx import Document
from PyPDF2 import PdfWriter

from backend.api.routes import router
from backend.services.ats_scorer import calculate_overall_score
from backend.services.resume_parser import validate_file, extract_text_from_docx


class BackendContractTests(unittest.TestCase):
    def test_groq_parser_module_exists_and_has_expected_functions(self):
        module = importlib.import_module("backend.services.groq_parser")
        self.assertTrue(hasattr(module, "parse_resume"))
        self.assertTrue(hasattr(module, "parse_job_description"))

    def test_shared_file_utils_exceptions_and_helpers_are_available(self):
        module = importlib.import_module("backend.utils.file_utils")
        for name in [
            "FileParsingError",
            "FileValidationError",
            "TextExtractionError",
            "FileUploadError",
            "log_error",
            "log_warning",
            "log_info",
            "with_fallback",
        ]:
            self.assertTrue(hasattr(module, name), f"Missing {name}")

    def test_validate_file_accepts_pdf_and_docx(self):
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        pdf_buffer = io.BytesIO()
        writer.write(pdf_buffer)
        pdf_bytes = pdf_buffer.getvalue()
        self.assertEqual(validate_file(pdf_bytes, "resume.pdf")[0], True)

        doc = Document()
        doc.add_paragraph("Python backend engineer")
        docbytes = io.BytesIO()
        doc.save(docbytes)
        self.assertEqual(validate_file(docbytes.getvalue(), "resume.docx")[0], True)

    def test_docx_extraction_returns_text(self):
        doc = Document()
        doc.add_paragraph("Python backend engineer")
        doc.add_paragraph("Built FastAPI services and dashboards")
        buffer = io.BytesIO()
        doc.save(buffer)
        text = extract_text_from_docx(buffer.getvalue())
        self.assertIn("Python backend engineer", text)
        self.assertIn("FastAPI", text)

    def test_health_route_returns_expected_shape(self):
        app = FastAPI()
        app.state.nlp = object()
        app.state.embedder = object()
        app.include_router(router)

        client = TestClient(app)
        response = client.get("/api/v1/health")

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["status"], "healthy")
        self.assertTrue(payload["nlp_loaded"])
        self.assertTrue(payload["embedder_loaded"])

    def test_ats_score_is_in_expected_range(self):
        score = calculate_overall_score(
            text="Senior Python engineer with strong backend and API experience. Built FastAPI services and deployed ML features.",
            parsed_resume={
                "experience": [{"job_title": "Engineer", "description": "Built APIs and dashboards", "duration_months": 18}],
                "education": [{"degree": "B.Tech", "institution": "College"}],
                "skills": ["Python", "FastAPI", "SQL", "AWS", "Docker"],
                "professional_summary": "Experienced Python engineer.",
                "projects": [{"title": "Analytics Portal", "description": "Built a FastAPI app for KPI dashboards"}],
            },
            skills=["Python", "FastAPI", "SQL", "AWS", "Docker"],
            keywords=["python", "fastapi", "sql", "aws", "docker", "api"],
            action_verbs=["Built", "Deployed", "Improved"],
            skill_validation_results={"validation_score": 12.0, "validation_percentage": 1.0},
            grammar_results={"penalty_applied": 0.0, "total_errors": 0},
            location_results={"penalty_applied": 0.0},
            jd_keywords=["python", "fastapi", "sql", "api"],
            experience_months=18,
        )

        self.assertGreaterEqual(score["overall_score"], 0.0)
        self.assertLessEqual(score["overall_score"], 100.0)
        self.assertGreaterEqual(score["formatting_score"], 0.0)
        self.assertLessEqual(score["formatting_score"], 20.0)

    def test_analyze_resume_route_uses_backend_contract(self):
        app = FastAPI()
        app.state.nlp = object()
        app.state.embedder = object()
        app.include_router(router)

        from backend.api import auth as auth_module
        app.dependency_overrides[auth_module.get_current_user] = lambda: "test-user"

        fake_result = {
            "ats_score": 88.0,
            "component_scores": {
                "formatting": 18.0,
                "keywords": 22.0,
                "content": 20.0,
                "skill_validation": 14.0,
                "ats_compatibility": 14.0,
            },
            "issues_summary": ["Clear summary"],
            "detailed_feedback": [],
            "jd_comparison": {
                "match_percentage": 84.0,
                "semantic_similarity": 0.85,
                "matched_keywords": ["python"],
                "missing_keywords": [],
                "skills_gap": [],
            },
            "skill_validation_details": {
                "validated": [{"skill": "Python", "projects": ["API project"]}],
                "unvalidated": [],
                "total": 1,
                "validated_count": 1,
                "validation_pct": 100.0,
            },
            "skills": ["Python"],
            "matched_keywords": ["python"],
            "missing_keywords": [],
            "interpretation": "Strong",
        }

        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        pdf_buffer = io.BytesIO()
        writer.write(pdf_buffer)
        pdf_bytes = pdf_buffer.getvalue()

        with patch("backend.services.resume_analyzer.analyze_full_resume", return_value=fake_result):
            client = TestClient(app)
            response = client.post(
                "/api/v1/analyze-resume",
                files={"resume": ("resume.pdf", pdf_bytes, "application/pdf")},
                data={"job_description": "Python backend engineer"},
            )

            self.assertEqual(response.status_code, 200)
            payload = response.json()
            self.assertEqual(payload["ATS_score"], 88.0)
            self.assertEqual(payload["ats_score"], 88.0)
            self.assertEqual(payload["keyword_match"], 84.0)


if __name__ == "__main__":
    unittest.main()
